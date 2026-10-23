"""Shared safe spreadsheet exports for web and offline use."""
import csv
import io
import json
import re
from fastapi.responses import Response

# XML 1.0 cannot represent these code points, although SQLite/JSON can retain
# them in pasted notes or source text. Encode the complete affected value rather
# than silently removing characters or letting openpyxl reject the export.
_XML_UNREPRESENTABLE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')

def safe_cell(value):
    text='' if value is None else str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) or text.startswith(('\t','\r','\n')) else text


JOB_HEADERS=['ID','Company','Title','Job family','Employment type','Advertised salary','Fit','Thai relevance','Status','Last verified','Source','Original URL','Description','Requirements','Preferred requirements','Why consider this','Requirements to check','Documented fit score /100','Score provisional','Score limits','Score version','Assessment summary','Description evidence date','Requirement evidence date','Field provenance','Clause evidence']
def job_row(j):
    provenance=j.get('field_provenance') or {}
    evidence_date=lambda field: (provenance.get(field) or {}).get('checked_at') or (provenance.get(field) or {}).get('source_record_date')
    return [j.get(k,'') for k in ('id','company','title','category')]+['; '.join(j.get('employment_types') or [j.get('employment_type') or 'Not stated'])]+[
        j.get('salary',{}).get('display','Not disclosed'),j.get('assessment',{}).get('fit'),j.get('thai_relevance'),j.get('status'),
        j.get('verification',{}).get('last_success_at'),'; '.join(x['name'] for x in j.get('sources',[])),j.get('url'),j.get('description'),
        '\n'.join(j.get('requirements',{}).get('mandatory') or []),'\n'.join(j.get('requirements',{}).get('preferred') or []),'\n'.join(j.get('assessment',{}).get('reasons') or []),'\n'.join(j.get('assessment',{}).get('gaps') or []),
        j.get('assessment',{}).get('score',{}).get('value'),j.get('assessment',{}).get('score',{}).get('provisional'),
        '\n'.join(x['reason'] for x in j.get('assessment',{}).get('score',{}).get('limits',[])), j.get('assessment',{}).get('score',{}).get('version'),
        j.get('assessment_summary'), evidence_date('description'), evidence_date('requirements'), json.dumps(provenance,ensure_ascii=False),
        json.dumps(j.get('assessment',{}).get('score',{}).get('duties',[]),ensure_ascii=False)]


def export_response(headers,rows,name,format='csv'):
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Font, Alignment
        book=Workbook();sheet=book.active;sheet.title='Roles'
        sheet.append(headers)
        full_text = None
        for sheet_row, row in enumerate(rows, 2):
            cells = []
            for index, value in enumerate(row):
                original = '' if value is None else str(value)
                encoded = bool(_XML_UNREPRESENTABLE.search(original))
                text = json.dumps(original, ensure_ascii=True) if encoded else original
                safe = safe_cell(text)
                if encoded or len(safe) > 32767:
                    # Excel's native cell limit is 32,767 characters. openpyxl
                    # silently truncates longer values, so retain full content
                    # in ordered, identified parts on a separate worksheet.
                    if full_text is None:
                        full_text = book.create_sheet('Full text')
                        full_text.append(['Record ID', 'Field', 'Part', 'Text', 'Formula prefix added',
                                          'Encoding', 'Main sheet row'])
                        full_text['F1'].comment = Comment(
                            'Group by Main sheet row and Field, order by Part, remove the first apostrophe '
                            'only when Formula prefix added is Yes, and join Text parts. For JSON string '
                            'encoding, decode that complete JSON string to recover every original character. '
                            'Plain text parts need no decoding.', 'SG Career Intelligence')
                        full_text.freeze_panes = 'A2'
                    # Split after JSON encoding: escaping can expand a short
                    # original beyond Excel's cell limit. Rejoin parts before
                    # decoding JSON, removing only labelled formula prefixes.
                    chunks = [text[start:start + 30000] for start in range(0, len(text), 30000)]
                    record_id = str(row[0])
                    if _XML_UNREPRESENTABLE.search(record_id):
                        record_id = json.dumps(record_id, ensure_ascii=True)
                    if len(record_id) > 32767:
                        record_id = record_id[:120] + f'… [ID in main sheet row {sheet_row}]'
                    for part, chunk in enumerate(chunks, 1):
                        escaped = safe_cell(chunk)
                        full_text.append([safe_cell(record_id), headers[index], part, escaped,
                                          'Yes' if escaped != chunk else 'No',
                                          'JSON string' if encoded else 'Plain text', sheet_row])
                    encoding_label = '; JSON string encoding' if encoded else ''
                    safe = safe_cell(text[:120] + f'… [Complete text: Full text worksheet; {len(chunks)} parts; field: {headers[index]}{encoding_label}]')
                cells.append(safe)
            sheet.append(cells)
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=28
        for row in sheet.iter_rows(min_row=2):
            for cell in row: cell.alignment=Alignment(vertical='top')
        if full_text is not None:
            for cell in full_text[1]: cell.font = Font(bold=True)
            for column in full_text.columns:
                full_text.column_dimensions[column[0].column_letter].width = 28
            full_text.column_dimensions['D'].width = 80
        data=io.BytesIO();book.save(data)
        return Response(data.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{name}.xlsx"'})
    data=io.StringIO();writer=csv.writer(data);writer.writerow(headers)
    writer.writerows([[safe_cell(v) for v in row] for row in rows])
    return Response(data.getvalue().encode('utf-8-sig'),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="{name}.csv"'})
