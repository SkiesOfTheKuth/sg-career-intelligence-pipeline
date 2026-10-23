"""API acceptance regressions for the complete workspace improvement."""
import csv
import io
import json
import unittest
import test_audit_qa as fixtures
from test_audit_qa import vacancy
from evidence import present_job, prepare_assessments


class WorkspaceV2API(unittest.TestCase):
    setUp = fixtures.WorkspaceAPI.setUp
    tearDown = fixtures.WorkspaceAPI.tearDown
    get = fixtures.WorkspaceAPI.get
    put = fixtures.WorkspaceAPI.put

    def test_employment_membership_options_and_legacy_filter(self):
        self.store.upsert_jobs([vacancy('mixed', employment_type='Full Time, Freelance'),
                                vacancy('unstated', employment_type=None)])
        ids = lambda q: [j['id'] for j in self.get('/api/jobs?' + q).json()['items']]
        self.assertIn('mixed', ids('opportunity_type=salaried'))
        self.assertIn('mixed', ids('opportunity_type=freelance'))
        self.assertIn('mixed', ids('employment_type=Full%20Time'))
        self.assertIn('mixed', ids('employment_type=Full-time'))
        self.assertIn('mixed', ids('employment_type=Full%20Time,%20Freelance'))
        self.assertEqual(ids('employment_type=Not%20stated'), ['unstated'])
        options = self.get('/api/options').json()['employment_type']
        self.assertIn('Full-time', options)
        self.assertIn('Freelance', options)
        self.assertIn('Not stated', options)
        self.assertFalse(any(',' in item for item in options))
        query = 'employment_type=Freelance&opportunity_type=all&sort_by=title'
        rows = list(csv.DictReader(io.StringIO(self.get('/api/export/filtered?' + query).content.decode('utf-8-sig'))))
        self.assertEqual([r['ID'] for r in rows], ids(query))

    def test_freshness_uses_success_not_attempt_or_historical_date(self):
        a = vacancy('successful', verified_at='2026-10-20T00:00:00+00:00')
        b = vacancy('historical', verified_at=None)
        b['verification'].update(checked_at='2099-01-01', last_success_at=None, status='historical')
        c = vacancy('failed-new-attempt', verified_at='2026-10-19T00:00:00+00:00')
        c['verification'].update(checked_at='2099-01-02', last_attempt_at='2099-01-02', status='failed')
        self.store.upsert_jobs([a,b,c])
        rows = self.get('/api/jobs?sort_by=freshness').json()['items']
        ids = [j['id'] for j in rows]
        self.assertLess(ids.index('successful'), ids.index('failed-new-attempt'))
        self.assertEqual(ids[-1], 'historical')
        exported = list(csv.DictReader(io.StringIO(self.get('/api/export/filtered?sort_by=freshness').content.decode('utf-8-sig'))))
        self.assertEqual([r['ID'] for r in exported], ids)

    def test_resolved_import_history_is_authenticated_paginated_and_exported(self):
        self.put('monthly', {'expected_revision':0, 'notes':'Server note'})
        self.client.post('/api/applications/import', auth=self.auth,
                         json={'records':{'monthly':{'status':'Applied', 'notes':'Original browser note'}}})
        current = self.store.get_application('monthly')
        conflict = current['import_conflicts'][0]
        response = self.put('monthly', {'expected_revision':current['revision'],
                                       'resolve_import_conflicts':[conflict['id']]})
        self.assertEqual(response.status_code, 200)
        path = '/api/applications/monthly/history'
        self.assertEqual(self.client.get(path).status_code, 401)
        history = self.get(path).json()
        self.assertIn('Original browser note', json.dumps(history))
        self.assertIn('Applied', json.dumps(history))
        self.assertIn('resolved_at', json.dumps(history))
        page = self.get(path + '?limit=1&offset=0').json()
        self.assertEqual(len(page['items']), 1)
        self.assertGreater(page['total'], 1)
        self.assertEqual(self.get(path + '?offset=-1').status_code, 422)
        self.assertEqual(self.get('/api/applications/missing/history').status_code, 404)
        exported = self.get('/api/export/applications?ids=monthly').content.decode('utf-8-sig')
        rows = list(csv.DictReader(io.StringIO(exported)))
        self.assertEqual(len(rows), 1)
        self.assertIn('Original browser note', exported)
        self.assertIn('resolved_at', exported)

    def test_selected_xlsx_retains_notes_and_history_as_text(self):
        from openpyxl import load_workbook
        self.put('monthly', {'expected_revision':0, 'notes':'=HYPERLINK("bad")'})
        self.put('annual', {'expected_revision':0, 'notes':'Unselected'})
        response = self.get('/api/export/applications?ids=monthly&format=xlsx')
        book = load_workbook(io.BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 1).value, 'monthly')
        self.assertEqual(sheet.cell(2, 5).data_type, 's')
        self.assertTrue(sheet.cell(2, 5).value.startswith("'="))

    def test_long_xlsx_text_is_complete_in_labelled_chunks_and_snapshot_exported(self):
        from openpyxl import load_workbook
        notes = '=SUM(1,2)' + 'long note ' * 4000
        self.put('monthly', {'expected_revision': 0, 'notes': notes})
        response = self.get('/api/export/applications?ids=monthly&format=xlsx')
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active.max_row, 2)
        self.assertIn('Job snapshot', [cell.value for cell in book.active[1]])
        self.assertIn('Full text', book.sheetnames)
        parts = [row for row in book['Full text'].iter_rows(min_row=2, values_only=True) if row[1] == 'Notes']
        restored = ''.join(row[3][1:] if row[4] == 'Yes' else row[3] for row in parts)
        self.assertEqual(restored, notes)
        self.assertTrue(all(cell.data_type != 'f' for row in book['Full text'] for cell in row))
        exported = list(csv.DictReader(io.StringIO(self.get('/api/export/applications?ids=monthly').content.decode('utf-8-sig'))))
        self.assertEqual(json.loads(exported[0]['Job snapshot'])['id'], 'monthly')

    def test_cached_presentation_tracks_changed_source_and_is_detached(self):
        job = vacancy('cache', requirements='Degree in Marketing')
        prepare_assessments([job])
        first = present_job(job)
        first['requirements']['mandatory'].append('Injected mutation')
        first['assessment']['score']['duties'].clear()
        again = present_job(job)
        self.assertNotIn('Injected mutation', again['requirements']['mandatory'])
        self.assertTrue(again['assessment']['score']['duties'])
        job['source_requirements'] = 'Japanese proficiency is required'
        changed = present_job(job)
        self.assertIn('Japanese proficiency is required', changed['requirements']['mandatory'])
        self.assertLessEqual(changed['assessment']['score']['value'],39)


if __name__ == '__main__':
    unittest.main()
